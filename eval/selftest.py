"""Self-test: read sample PDFs with the real engine and score the final rows - what reaches the
Excel - against a truth table, per field.

Every field lands in one of five bins:

* ``green_right`` - shown as verified and right;
* ``GREEN_WRONG`` - shown as verified and WRONG: a silent error, the number that must be zero;
* ``orange_right`` - flagged for review, but the value shown is right;
* ``orange_other`` - flagged, and the value is wrong or missing;
* ``truth_unknown`` - the truth table is empty or says "לא ברור" (draft cells are not scored).

    python -m eval.selftest <truth.xlsx> <part1.pdf> <part2.pdf> --live [--pages 1,2,5]

The truth workbook has one sheet per PDF ("חלק 1", "חלק 2") with a "עמוד בקובץ" column. A
live run SENDS EVERY PAGE TO THE CLOUD MODEL and spends paid credit (about three requests a page),
so it does nothing without ``--live``. The scoring itself is pure and tested offline.
"""

from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from collections.abc import Iterable, Sequence
from pathlib import Path
import re
import sys
from typing import Final, Literal

from backend.core.models import CommitmentRow, ConfidenceLevel, FieldKey

COLUMNS: Final = {
    FieldKey.first_name: "שם פרטי",
    FieldKey.last_name: "שם משפחה",
    FieldKey.national_id: "ת.ז.",
    FieldKey.central_number: "מספר מרכזי",
    FieldKey.receipt_date: "תאריך קבלה",
    FieldKey.sender: "גורם שולח",
    FieldKey.commitment_number: "מספר התחייבות",
}
NAMES: Final = (FieldKey.first_name, FieldKey.last_name)
BINS: Final = ("green_right", "GREEN_WRONG", "orange_right", "orange_other", "truth_unknown")
Bin = Literal["green_right", "GREEN_WRONG", "orange_right", "orange_other", "truth_unknown"]

# One institution under several names: the draft truth table and the closed list differ.
_ALIASES: Final = {"שמיר": "אסףהרופא", "המרכזהלאומילרפואהמשפטית": "המכוןלרפואהמשפטית"}
_GENERIC: Final = ("שירותיבריאות", "מעבדת", "מרכזרפואי", "ביתחולים")
_UNKNOWN: Final = "לא ברור"


def sender_key(text: str) -> str:
    """A comparable institution: "כללית (מרכז רפואי כרמל)" -> "כרמל" (the one in brackets),
    "מכבי שירותי בריאות" -> "מכבי"."""
    inner = re.search(r"\((.*?)\)", text)
    word = re.sub(r"[^֐-׿]", "", inner[1] if inner else text)
    for generic in _GENERIC:
        word = word.replace(generic, "")
    return _ALIASES.get(word, word)


def tokens(key: FieldKey, value: object) -> frozenset[str]:
    """The comparable form of a cell: dates as DD/MM/YYYY, institutions by `sender_key`, names
    by their letters, numbers as alphanumeric tokens without leading zeros ("947468; 947469")."""
    text = str(value or "").strip()
    if not text:
        return frozenset()
    if key is FieldKey.receipt_date:
        m = re.search(r"(\d{1,2})\D+(\d{1,2})\D+(\d{2,4})", text)
        if not m:
            return frozenset({text})
        year = int(m[3]) + (2000 if int(m[3]) < 100 else 0)
        return frozenset({f"{int(m[1]):02d}/{int(m[2]):02d}/{year}"})
    if key is FieldKey.sender:
        return frozenset({sender_key(text)})
    if key in NAMES:
        return frozenset({re.sub(r"[^֐-׿A-Za-z]", "", text)})
    return frozenset(t.lstrip("0") or "0" for t in re.findall(r"[A-Za-z0-9]+", text))


def classify(key: FieldKey, truth: object, shown: Iterable[str], green: bool) -> Bin:
    """The bin of one field: what the truth table says, what the rows show, and whether every
    row shows it as verified."""
    truth_text = str(truth or "")
    expected = tokens(key, truth_text)
    if not expected or _UNKNOWN in truth_text:
        return "truth_unknown"
    got = frozenset().union(*(tokens(key, v) for v in shown))
    if got == expected:
        return "green_right" if green else "orange_right"
    return "GREEN_WRONG" if green else "orange_other"


def field_of(rows: Sequence[CommitmentRow], key: FieldKey) -> tuple[list[str], bool]:
    """(the values the rows of one page show for `key`, True when every row shows it green).
    Central numbers come from every row's list: a form may carry several."""
    values: list[str] = []
    green = bool(rows)
    for row in rows:
        result = row.fields.get(key)
        if key is FieldKey.central_number and row.central_numbers:
            values += row.central_numbers
        elif result is not None and result.value:
            values.append(result.value)
        if result is None or result.level is ConfidenceLevel.D:
            green = False
    return list(dict.fromkeys(values)), green


def score(
    rows_by_page: dict[tuple[int, int], list[CommitmentRow]],
    truth: dict[tuple[int, int], dict[str, object]],
) -> tuple[dict[FieldKey, Counter[str]], list[str]]:
    """({field: bin counts}, a note per non-green-right field). Names are counted only where the
    truth table has them. Pure: no engine, no file."""
    stats: dict[FieldKey, Counter[str]] = defaultdict(Counter)
    notes: list[str] = []
    for (part, page), rows in sorted(rows_by_page.items()):
        row = truth.get((part, page))
        if row is None:
            continue
        for key, column in COLUMNS.items():
            shown, green = field_of(rows, key)
            got = classify(key, row.get(column), shown, green)
            if key in NAMES and got == "truth_unknown":
                continue
            stats[key][got] += 1
            if got != "green_right" and got != "truth_unknown":
                notes.append(
                    f"{got:13} p{part}.{page} {key.value}: {shown!r} vs {row.get(column)!r}"
                )
    return stats, notes


def truth_rows(path: Path) -> dict[tuple[int, int], dict[str, object]]:
    """{(part, page): truth row} from a workbook with one sheet per part."""
    import openpyxl

    rows: dict[tuple[int, int], dict[str, object]] = {}
    wb = openpyxl.load_workbook(path, read_only=True)
    for part, title in ((1, "חלק 1"), (2, "חלק 2")):
        if title not in wb.sheetnames:
            continue
        it = wb[title].iter_rows(values_only=True)
        header = [str(h) for h in next(it)]
        for values in it:
            row = dict(zip(header, values, strict=False))
            page = row.get("עמוד בקובץ")
            if isinstance(page, int | float | str) and str(page).strip().isdigit():
                rows[(part, int(str(page).strip()))] = row
    return rows


def read_rows(pdfs: Sequence[Path], only: set[int]) -> dict[tuple[int, int], list[CommitmentRow]]:
    """Run the real engine (cloud model included) over the PDFs and group the rows by page."""
    from backend.container import build_container
    from backend.core.pipeline.orchestrator import iter_outcomes
    from backend.core.pipeline.s0_intake import UploadedBlob, intake
    from backend.core.pipeline.s10_rows import RowContext, assemble_rows
    from backend.settings import Settings

    svc = build_container(Settings())  # type: ignore[call-arg]
    ctx = RowContext(
        clock=svc.clock,
        senders=svc.senders,
        weights=svc.pipeline.weights,
        scoring=svc.pipeline.scoring,
        config=svc.pipeline.rows,
    )
    by_page: dict[tuple[int, int], list[CommitmentRow]] = defaultdict(list)
    for part, pdf in enumerate(pdfs, start=1):
        result = intake(
            [UploadedBlob(filename=pdf.name, data=pdf.read_bytes())],
            settings=svc.settings,
            documents=svc.documents,
            storage=svc.storage,
            job_id=f"selftest{part}",
            known_hashes=frozenset(),
        )
        pages = [p for p in result.pages if not only or p.page_index + 1 in only]
        try:
            outcomes = list(iter_outcomes(pages, services=svc, cloud_consent=True))
            rows, _, _ = assemble_rows(outcomes, ctx=ctx)
        finally:
            svc.storage.purge_job(f"selftest{part}")
        for row in rows:
            by_page[(part, int(row.page_ids[0].rsplit("#", 1)[-1]))].append(row)
    return by_page


def report(stats: dict[FieldKey, Counter[str]], notes: Sequence[str]) -> str:
    lines = [f"{'field':20}" + "".join(f"{b:>15}" for b in BINS)]
    total: Counter[str] = Counter()
    for key in COLUMNS:
        if key not in stats:
            continue
        total.update(stats[key])
        lines.append(f"{key.value:20}" + "".join(f"{stats[key][b]:>15}" for b in BINS))
    lines.append(f"{'TOTAL':20}" + "".join(f"{total[b]:>15}" for b in BINS))
    return "\n".join([*lines, "", *notes])


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0] if __doc__ else None)
    ap.add_argument("truth", type=Path)
    ap.add_argument("pdfs", type=Path, nargs="+")
    ap.add_argument("--pages", default="", help="comma-separated 1-based pages, e.g. 1,2,5")
    ap.add_argument("--live", action="store_true", help="really read the pages (spends credit)")
    args = ap.parse_args(argv)
    if not args.live:
        print("A self-test reads every page with the cloud model and spends credit. Add --live.")
        return 2
    only = {int(p) for p in args.pages.split(",") if p.strip()}
    stats, notes = score(read_rows(args.pdfs, only), truth_rows(args.truth))
    print(report(stats, notes))
    return 1 if sum(s["GREEN_WRONG"] for s in stats.values()) else 0


if __name__ == "__main__":
    sys.exit(main())
